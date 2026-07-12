#!/usr/bin/env python3
"""Build the Caisson attorney/CPA DOCX package from reviewed Markdown drafts."""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor, Twips

TABLE_HELPER_DIR = Path(
    "/Users/liamt/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.709.11516/skills/documents/scripts"
)
sys.path.insert(0, str(TABLE_HELPER_DIR))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


ROOT = Path("/Users/liamt/files/projects/caisson/recovered-codex-outputs")
DRAFT_DIR = ROOT / "work" / "drafts"
BUILD_DIR = ROOT / "work" / "build_docs" / "unscrubbed"

BANNER = "DRAFT FOR DISCUSSION - NOT FOR SIGNATURE OR FILING - ATTORNEY/CPA REVIEW REQUIRED"
NAVY = "18324B"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5D6874"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
PALE_BLUE = "EEF5FA"
PALE_GOLD = "FFF5DA"
PALE_RED = "FCE8E6"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111111"
RULE = "CBD5DF"


@dataclass(frozen=True)
class DocConfig:
    source: Path
    output_name: str
    kicker: str
    title: str
    subtitle: str
    prepared_for: str
    status: str
    preset: str
    accent: str
    instrument_mode: bool = False
    compact_memo: bool = False


CONFIGS = [
    DocConfig(
        source=DRAFT_DIR / "counsel-drafting-memorandum.md",
        output_name="caisson-counsel-drafting-memorandum-DRAFT-2026-07-11.docx",
        kicker="COUNSEL DRAFTING MEMORANDUM",
        title="Caisson Software LLC",
        subtitle="Present ownership, trust structure, governance, succession, IP, and employment",
        prepared_for="Vik Thadani's counsel - conflict and engagement-scope review",
        status="Attorney redline requested",
        preset="standard_business_brief",
        accent=BLUE,
        compact_memo=True,
    ),
    DocConfig(
        source=DRAFT_DIR / "draft-instrument-set.md",
        output_name="caisson-draft-instrument-set-FOR-COUNSEL-2026-07-11.docx",
        kicker="ATTORNEY-REDLINE INSTRUMENT SET",
        title="Caisson Ownership and Control Drafts",
        subtitle="Trust, assignment, operating agreement, succession, IP, employment, and age-18 closing documents",
        prepared_for="Georgia counsel - drafting and enforceability review",
        status="Substantially complete discussion drafts",
        preset="contract_negotiation_brief",
        accent=DARK_BLUE,
        instrument_mode=True,
    ),
    DocConfig(
        source=DRAFT_DIR / "cpa-coordination-packet.md",
        output_name="caisson-cpa-coordination-packet-DRAFT-2026-07-11.docx",
        kicker="CPA COORDINATION PACKET",
        title="Caisson Finance and Tax Coordination",
        subtitle="Working reconstruction, family funding, trust sequencing, payroll, and opening-entry decisions",
        prepared_for="Vik Thadani's CPA - existing-client advisory intake",
        status="CPA review requested; not an approved ledger",
        preset="standard_business_brief",
        accent="1D7169",
    ),
]


def ascii_text(text: str) -> str:
    replacements = {
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u00a0": " ",
        "\u2026": "...",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def set_run_font(run, *, name: str = "Calibri", size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None, underline: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table, color: str = "CDD4DC", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, side: str, color: str, size: str = "8",
                         space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), space)
    edge.set(qn("w:color"), color)


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8, color=MUTED)


def add_hyperlink(paragraph, label: str, url: str, *, size: float | None = None) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    run_properties.extend([fonts, color, underline])
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        run_properties.append(sz)
    text = OxmlElement("w:t")
    text.text = ascii_text(label)
    run.extend([run_properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(
    r"(\[[^\]]+\]\(https?://[^)]+\)|<https?://[^>]+>|https?://[^\s)>]+|\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
)


def add_inline(paragraph, text: str, *, size: float | None = None,
               color: str | None = None, bold: bool = False) -> None:
    text = ascii_text(text)
    position = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2), size=size)
        elif token.startswith("<http://") or token.startswith("<https://"):
            url = token[1:-1]
            add_hyperlink(paragraph, url, url, size=size)
        elif token.startswith("http://") or token.startswith("https://"):
            trailing = ""
            while token and token[-1] in ".,;:":
                trailing = token[-1] + trailing
                token = token[:-1]
            add_hyperlink(paragraph, token, token, size=size)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, size=size, color=color, bold=bold)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=(size or 10) - 0.5, color=DARK_BLUE)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color, bold=bold)


def add_numbering_definition(document: Document, *, ordered: bool, compact: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(element.get(qn("w:abstractNumId"))) for element in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    marker = 270 if compact else 360
    text_indent = 540 if compact else 720
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), f"%{level + 1}." if ordered else ("•" if level == 0 else "-"))
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(text_indent + level * 360))
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(text_indent + level * 360))
        indent.set(qn("w:hanging"), str(text_indent - marker))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80" if compact else "160")
        spacing.set(qn("w:line"), "300" if compact else "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, indent, spacing])
        lvl.extend([start, num_fmt, lvl_text, justification, p_pr])
        if not ordered:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Calibri")
            r_fonts.set(qn("w:hAnsi"), "Calibri")
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(min(level, 2)))
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])


def configure_styles(document: Document, config: DocConfig) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5 if config.instrument_mode else (10.25 if config.compact_memo else 11))
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15 if config.instrument_mode else (1.04 if config.compact_memo else 1.10)
    normal.paragraph_format.widow_control = True

    heading_values = {
        "Heading 1": (16, config.accent, 14 if config.instrument_mode else 16, 8),
        "Heading 2": (13, config.accent, 11 if config.instrument_mode else 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_values.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Table Text" not in [style.name for style in styles]:
        style = styles.add_style("Table Text", 1)
    table_style = styles["Table Text"]
    table_style.font.name = "Calibri"
    table_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_style.font.size = Pt(9 if config.compact_memo else (9.25 if not config.instrument_mode else 9))
    table_style.paragraph_format.space_after = Pt(2)
    table_style.paragraph_format.line_spacing = 1.05


def configure_page(document: Document, config: DocConfig) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    vertical_margin = 0.78 if config.compact_memo else 1
    horizontal_margin = 0.74 if config.compact_memo else 1
    section.top_margin = Inches(vertical_margin)
    section.right_margin = Inches(horizontal_margin)
    section.bottom_margin = Inches(vertical_margin)
    section.left_margin = Inches(horizontal_margin)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def build_header_footer(document: Document, config: DocConfig) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(1)
    left = paragraph.add_run("CAISSON SOFTWARE LLC")
    set_run_font(left, size=8, color=NAVY, bold=True)
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    paragraph.add_run("\t")
    right = paragraph.add_run("DRAFT FOR DISCUSSION")
    set_run_font(right, size=8, color=RED, bold=True)
    set_paragraph_border(paragraph, side="bottom", color=RULE, size="4", space="3")

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(1)
    set_paragraph_border(paragraph, side="top", color=RULE, size="4", space="3")
    left = paragraph.add_run("CONFIDENTIAL - NOT FOR SIGNATURE OR FILING")
    set_run_font(left, size=8, color=MUTED)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    paragraph.add_run("\tPage ")
    add_field(paragraph, "PAGE")
    tail = paragraph.add_run(" of ")
    set_run_font(tail, size=8, color=MUTED)
    add_field(paragraph, "NUMPAGES")


def add_banner(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Pt(6)
    paragraph.paragraph_format.right_indent = Pt(6)
    set_paragraph_shading(paragraph, RED)
    run = paragraph.add_run(BANNER)
    set_run_font(run, size=8.5, color=WHITE, bold=True)


def add_metadata_table(document: Document, config: DocConfig) -> None:
    rows = [
        ("Prepared for", config.prepared_for),
        ("Prepared", "July 11, 2026"),
        ("Status", config.status),
        ("Governing baseline", "Georgia law; counsel to confirm trust jurisdiction and family-law effects"),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    widths = [1900, 7460]
    for row_index, (label, value) in enumerate(rows):
        set_row_cant_split(table.rows[row_index])
        label_cell, value_cell = table.rows[row_index].cells
        set_cell_shading(label_cell, LIGHT_GRAY)
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.style = document.styles["Table Text"]
        label_run = label_paragraph.add_run(label)
        set_run_font(label_run, size=9, color=NAVY, bold=True)
        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.style = document.styles["Table Text"]
        add_inline(value_paragraph, value, size=9.2)
    set_table_borders(table, color="D7DDE4", size="4")
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_title_block(document: Document, config: DocConfig) -> None:
    add_banner(document)
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(4)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run(config.kicker)
    set_run_font(run, size=9.5, color=config.accent, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(config.title)
    set_run_font(run, size=25 if not config.instrument_mode else 23, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(config.subtitle)
    set_run_font(run, size=12.5, color=MUTED)
    add_metadata_table(document, config)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(2)
    rule.paragraph_format.space_after = Pt(8)
    set_paragraph_border(rule, side="bottom", color=config.accent, size="10", space="2")


def markdown_table_widths(rows: list[list[str]]) -> list[int]:
    count = len(rows[0])
    if count == 2:
        first_mean = sum(len(row[0]) for row in rows) / len(rows)
        return [2100, 7260] if first_mean < 28 else [3300, 6060]
    if count == 3:
        return [1800, 3400, 4160]
    if count == 4:
        return [1300, 2500, 2760, 2800]
    headers = [cell.strip().lower() for cell in rows[0]]
    if count == 5 and headers[-1] == "payer":
        return [2000, 3500, 1800, 1100, 960]
    if count == 6 and headers[2] == "payer":
        return [2000, 1800, 850, 1200, 2200, 1310]
    if count == 7 and rows[0][-1].strip().lower() == "evidence owner":
        return [1450, 1400, 1400, 1400, 1400, 1400, 910]
    if count == 7 and "commit/file range" in headers:
        return [1550, 1250, 1350, 1150, 1650, 1050, 1360]
    if count == 7 and rows[0][-1].strip().lower() == "compliance status":
        return [1300, 1050, 1050, 1300, 2150, 1360, 1150]
    weights = []
    for column in range(count):
        longest = max(len(row[column]) for row in rows)
        weights.append(max(8, min(45, longest)))
    return column_widths_from_weights(weights, 9360)


def add_single_markdown_table(document: Document, rows: list[list[str]], config: DocConfig) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    widths = markdown_table_widths(rows)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER if row_index == 0 else WD_ALIGN_VERTICAL.TOP
            if row_index == 0:
                set_cell_shading(cell, config.accent if config.instrument_mode else LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Text"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if row_index == 0:
                add_inline(
                    paragraph,
                    value,
                    size=8.8,
                    color=WHITE if config.instrument_mode else NAVY,
                    bold=True,
                )
            else:
                add_inline(paragraph, value, size=9 if config.instrument_mode else 9.2)
    set_table_borders(table, color="C8D0D9", size="4")
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_table(document: Document, rows: list[list[str]], config: DocConfig) -> None:
    """Add a readable table, splitting very wide instrument schedules by field group."""
    if not rows:
        return
    column_count = len(rows[0])
    if config.instrument_mode and column_count > 7:
        remaining = list(range(1, column_count))
        split_at = (len(remaining) + 1) // 2
        groups = [remaining[:split_at], remaining[split_at:]]
        for group_index, group in enumerate(groups):
            if group_index:
                note = document.add_paragraph()
                note.paragraph_format.keep_with_next = True
                note.paragraph_format.space_before = Pt(2)
                note.paragraph_format.space_after = Pt(4)
                add_inline(
                    note,
                    "*Continuation - remaining fields for the same rows.*",
                    size=8.8,
                    color=MUTED,
                )
            indices = [0, *group]
            split_rows = [[row[index] for index in indices] for row in rows]
            add_single_markdown_table(document, split_rows, config)
        return
    add_single_markdown_table(document, rows, config)


def add_callout(document: Document, text: str, config: DocConfig, *, caution: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.08
    set_paragraph_shading(paragraph, PALE_GOLD if caution else PALE_BLUE)
    set_paragraph_border(paragraph, side="left", color="A66F00" if caution else config.accent, size="18", space="6")
    add_inline(paragraph, text, size=9.8 if config.instrument_mode else 10)


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return lines[index].strip().startswith("|") and bool(
        re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1])
    )


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    width = len(rows[0])
    rows = [row + [""] * (width - len(row)) for row in rows]
    return rows, index


def is_special(line: str, lines: list[str], index: int) -> bool:
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith(">")
        or stripped.startswith("```")
        or bool(re.match(r"^\s*[-*]\s+", line))
        or bool(re.match(r"^\s*\d+\.\s+", line))
        or stripped in {"---", "***"}
        or is_table_start(lines, index)
    )


def add_body_paragraph(document: Document, text: str, config: DocConfig) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = True
    size = 10.5 if config.instrument_mode else (10.25 if config.compact_memo else 11)
    add_inline(paragraph, text, size=size)


def add_heading(document: Document, level: int, text: str, config: DocConfig) -> None:
    mapped = max(1, min(3, level - 1 if level > 1 else 1))
    should_break = config.instrument_mode and (
        level == 1
        or (mapped == 1 and text.lower().startswith("schedule"))
    ) and len(document.paragraphs) > 8
    paragraph = document.add_paragraph(style=f"Heading {mapped}")
    paragraph.paragraph_format.page_break_before = should_break
    add_inline(paragraph, text, size=None, color=None, bold=True)


def render_markdown(document: Document, markdown: str, config: DocConfig) -> None:
    markdown = ascii_text(markdown)
    lines = markdown.splitlines()
    if lines and lines[0].strip() == "---":
        closing = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), None)
        if closing is not None:
            lines = lines[closing + 1 :]
    cleaned: list[str] = []
    source_titles = {
        "Caisson Software LLC - Counsel Drafting Memorandum",
        "Caisson Software LLC - Tax, Trust, Payroll, and Opening-Books Coordination Packet",
        "Caisson Software LLC - Draft Integrated Instrument Set",
    }
    for line in lines:
        plain = line.strip().lstrip(">").strip()
        if plain == BANNER:
            continue
        heading = re.match(r"^#{1,2}\s+(.+)$", line.strip())
        if heading and heading.group(1).strip() in source_titles:
            continue
        cleaned.append(line.rstrip())
    lines = cleaned

    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            set_paragraph_shading(paragraph, LIGHT_GRAY)
            run = paragraph.add_run(stripped)
            set_run_font(run, name="Consolas", size=9, color=DARK_BLUE)
            index += 1
            continue
        if is_table_start(lines, index):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows, config)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            add_heading(document, len(heading.group(1)), heading.group(2), config)
            index += 1
            continue
        if stripped in {"---", "***"}:
            if not config.instrument_mode:
                rule = document.add_paragraph()
                set_paragraph_border(rule, side="bottom", color=RULE, size="4", space="2")
                rule.paragraph_format.space_after = Pt(6)
            index += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            text = " ".join(quote_lines)
            add_callout(document, text, config, caution=bool(re.search(r"warning|risk|do not|not for", text, re.I)))
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        ordered = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if bullet or ordered:
            is_ordered = ordered is not None
            list_items: list[tuple[int, str]] = []
            matcher = ordered if ordered else bullet
            assert matcher is not None
            while index < len(lines):
                current = re.match(r"^(\s*)\d+\.\s+(.+)$", lines[index]) if is_ordered else re.match(r"^(\s*)[-*]\s+(.+)$", lines[index])
                if not current:
                    break
                level = min(len(current.group(1)) // 2, 2)
                list_items.append((level, current.group(2).strip()))
                index += 1
            num_id = add_numbering_definition(document, ordered=is_ordered, compact=config.instrument_mode)
            for level, text in list_items:
                paragraph = document.add_paragraph()
                apply_numbering(paragraph, num_id, level)
                paragraph.paragraph_format.widow_control = True
                if config.compact_memo:
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    paragraph.paragraph_format.line_spacing = 1.0
                size = 10.4 if config.instrument_mode else (9.95 if config.compact_memo else 10.8)
                add_inline(paragraph, text, size=size)
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and not is_special(lines[index], lines, index):
            paragraph_lines.append(lines[index].strip())
            index += 1
        add_body_paragraph(document, " ".join(paragraph_lines), config)


def build(config: DocConfig) -> Path:
    if not config.source.exists():
        raise FileNotFoundError(config.source)
    document = Document()
    configure_page(document, config)
    configure_styles(document, config)
    build_header_footer(document, config)
    add_title_block(document, config)
    render_markdown(document, config.source.read_text(encoding="utf-8"), config)

    properties = document.core_properties
    properties.title = config.title
    properties.subject = config.subtitle
    properties.category = "Confidential attorney/CPA discussion draft"
    properties.keywords = "Caisson Software LLC; draft; attorney review; CPA review"

    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    output = BUILD_DIR / config.output_name
    document.save(output)
    return output


def main() -> None:
    outputs = [build(config) for config in CONFIGS]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
